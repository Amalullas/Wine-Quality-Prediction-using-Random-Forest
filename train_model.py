"""
Wine Quality Prediction - Assignment 2 (Machine Learning Application)

Run:  python train_model.py

What it does (in order):
  1. Loads the Wine Quality (red) dataset
  2. Explores the data and saves EDA plots
  3. Preprocesses it (duplicates, missing values, train/test split, scaling)
  4. Trains Linear Regression and Random Forest
  5. Evaluates both (R2, MAE, RMSE, accuracy-style metrics, cross-validation)
  6. Saves result plots, the trained model, and metrics
  7. Auto-generates the Word report (outputs/Wine_Quality_Report.docx)
"""

import datetime
import json
import os
import urllib.request
import warnings

import joblib
import matplotlib

matplotlib.use("Agg")  # no pop-up windows, just save images
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from sklearn.ensemble import RandomForestRegressor
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.model_selection import cross_val_score, train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

warnings.filterwarnings("ignore")

# ------------------------------------------------------------------
# EDIT THESE FOUR LINES - they go on the first page of the report
# ------------------------------------------------------------------
STUDENT_NAME = "Amal"
ROLL_NUMBER = "6"
COLLEGE = "Universal Engineering College, Thrissur"
COURSE = "Machine Learning"
# ------------------------------------------------------------------

DATA_PATH = "data/winequality-red.csv"
DATA_URL = (
    "https://archive.ics.uci.edu/ml/machine-learning-databases/"
    "wine-quality/winequality-red.csv"
)
OUT = "outputs"
RANDOM_STATE = 42

FEATURE_DESCRIPTIONS = {
    "fixed acidity": "Non-volatile acids (mainly tartaric acid), g/dm3",
    "volatile acidity": "Acetic acid content; too much gives a vinegar taste, g/dm3",
    "citric acid": "Adds freshness and flavour, g/dm3",
    "residual sugar": "Sugar left after fermentation, g/dm3",
    "chlorides": "Salt content, g/dm3",
    "free sulfur dioxide": "Free SO2, protects wine from spoilage, mg/dm3",
    "total sulfur dioxide": "Free + bound SO2, mg/dm3",
    "density": "Density of the wine, g/cm3",
    "pH": "Acidity level on the pH scale",
    "sulphates": "Potassium sulphate, an additive/preservative, g/dm3",
    "alcohol": "Alcohol content, % by volume",
    "quality": "TARGET - expert quality score (0 to 10; observed range 3 to 8)",
}


# ------------------------------------------------------------------
# 1. LOAD DATA
# ------------------------------------------------------------------
def load_data():
    os.makedirs("data", exist_ok=True)
    if not os.path.exists(DATA_PATH):
        print("Dataset not found locally - downloading from UCI ...")
        urllib.request.urlretrieve(DATA_URL, DATA_PATH)
    # The UCI file uses ';' as separator, some mirrors use ','. Detect it.
    with open(DATA_PATH, encoding="utf-8-sig") as f:
        sep = ";" if ";" in f.readline() else ","
    df = pd.read_csv(DATA_PATH, sep=sep, encoding="utf-8-sig")
    df.columns = [c.strip().strip('"') for c in df.columns]
    return df


# ------------------------------------------------------------------
# REPORT HELPERS
# ------------------------------------------------------------------
def add_par(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_fig(doc, path, caption, width=5.8):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True


def add_table(doc, header, rows):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    doc.add_paragraph()


# ------------------------------------------------------------------
# METRICS
# ------------------------------------------------------------------
def evaluate(y_true, y_pred):
    y_true = np.asarray(y_true)
    y_pred = np.asarray(y_pred)
    rounded = np.clip(np.round(y_pred), 3, 8)
    return {
        "R2": float(r2_score(y_true, y_pred)),
        "MAE": float(mean_absolute_error(y_true, y_pred)),
        "RMSE": float(np.sqrt(mean_squared_error(y_true, y_pred))),
        "Exact_Accuracy": float((rounded == y_true).mean()),
        "Within_1_Point": float((np.abs(rounded - y_true) <= 1).mean()),
    }


def main():
    os.makedirs(OUT, exist_ok=True)
    sns.set_theme(style="whitegrid")

    # --------------------------------------------------------------
    # 1-2. LOAD + EXPLORE
    # --------------------------------------------------------------
    df = load_data()
    n_rows_raw, n_cols = df.shape
    n_missing = int(df.isnull().sum().sum())
    n_dupes = int(df.duplicated().sum())
    print(f"Raw data: {n_rows_raw} rows x {n_cols} columns")
    print(f"Missing values: {n_missing} | Duplicate rows: {n_dupes}")

    # Plot 1: distribution of the target
    plt.figure(figsize=(7, 4.5))
    ax = sns.countplot(x="quality", data=df, color="#7b2d43")
    for c in ax.containers:
        ax.bar_label(c)
    plt.title("Distribution of Wine Quality Scores")
    plt.xlabel("Quality score")
    plt.ylabel("Number of wines")
    plt.tight_layout()
    plt.savefig(f"{OUT}/01_quality_distribution.png", dpi=150)
    plt.close()

    # Plot 2: correlation heatmap
    plt.figure(figsize=(10, 8))
    sns.heatmap(df.corr(), annot=True, fmt=".2f", cmap="coolwarm", center=0,
                annot_kws={"size": 8})
    plt.title("Correlation Heatmap")
    plt.tight_layout()
    plt.savefig(f"{OUT}/02_correlation_heatmap.png", dpi=150)
    plt.close()

    # Plot 3: alcohol vs quality
    plt.figure(figsize=(7, 4.5))
    sns.boxplot(x="quality", y="alcohol", data=df, color="#c98b9b")
    plt.title("Alcohol Content vs Wine Quality")
    plt.tight_layout()
    plt.savefig(f"{OUT}/03_alcohol_vs_quality.png", dpi=150)
    plt.close()

    # --------------------------------------------------------------
    # 3. PREPROCESSING
    # --------------------------------------------------------------
    if n_missing > 0:
        df = df.fillna(df.median(numeric_only=True))
    df = df.drop_duplicates().reset_index(drop=True)
    n_rows = len(df)
    print(f"After removing duplicates: {n_rows} rows")

    X = df.drop(columns="quality")
    y = df["quality"]
    feature_names = list(X.columns)

    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=0.2, random_state=RANDOM_STATE, stratify=y
    )
    print(f"Train: {len(X_train)} rows | Test: {len(X_test)} rows")

    # --------------------------------------------------------------
    # 4-5. TRAIN + TEST BOTH MODELS
    # (scaling lives inside a Pipeline so the test set never leaks in)
    # --------------------------------------------------------------
    models = {
        "Linear Regression": Pipeline([
            ("scaler", StandardScaler()),
            ("model", LinearRegression()),
        ]),
        "Random Forest": Pipeline([
            ("scaler", StandardScaler()),
            ("model", RandomForestRegressor(
                n_estimators=300, random_state=RANDOM_STATE, n_jobs=-1)),
        ]),
    }

    results, predictions = {}, {}
    for name, pipe in models.items():
        cv = cross_val_score(pipe, X_train, y_train, cv=5, scoring="r2")
        pipe.fit(X_train, y_train)
        pred = pipe.predict(X_test)
        predictions[name] = pred
        results[name] = evaluate(y_test, pred)
        results[name]["CV_R2"] = float(cv.mean())
        results[name]["CV_R2_std"] = float(cv.std())
        r = results[name]
        print(f"\n{name}\n  R2={r['R2']:.3f}  MAE={r['MAE']:.3f}  RMSE={r['RMSE']:.3f}"
              f"  CV R2={r['CV_R2']:.3f}")

    # Pick the winner using cross-validation on the TRAINING data only
    best_name = max(results, key=lambda n: results[n]["CV_R2"])
    best_pipe = models[best_name]
    best_pred = predictions[best_name]
    best = results[best_name]
    print(f"\nBest model: {best_name}")

    # --------------------------------------------------------------
    # 6. RESULT PLOTS
    # --------------------------------------------------------------
    # Plot 4: actual vs predicted
    rng = np.random.default_rng(RANDOM_STATE)
    jitter = rng.uniform(-0.12, 0.12, size=len(y_test))
    plt.figure(figsize=(6.5, 6))
    plt.scatter(np.asarray(y_test) + jitter, best_pred, alpha=0.55, color="#7b2d43")
    plt.plot([3, 8], [3, 8], "k--", label="Perfect prediction")
    plt.xlabel("Actual quality")
    plt.ylabel("Predicted quality")
    plt.title(f"Actual vs Predicted ({best_name})")
    plt.legend()
    plt.tight_layout()
    plt.savefig(f"{OUT}/04_actual_vs_predicted.png", dpi=150)
    plt.close()

    # Plot 5: residuals
    residuals = np.asarray(y_test) - best_pred
    plt.figure(figsize=(7, 4.5))
    sns.histplot(residuals, bins=25, kde=True, color="#7b2d43")
    plt.axvline(0, color="k", linestyle="--")
    plt.title("Residuals (Actual - Predicted)")
    plt.xlabel("Error in quality points")
    plt.tight_layout()
    plt.savefig(f"{OUT}/05_residuals.png", dpi=150)
    plt.close()

    # Plot 6: feature importance
    inner = best_pipe.named_steps["model"]
    if hasattr(inner, "feature_importances_"):
        importance = inner.feature_importances_
    else:
        importance = np.abs(inner.coef_) / np.abs(inner.coef_).sum()
    imp = pd.Series(importance, index=feature_names).sort_values()
    plt.figure(figsize=(7.5, 5))
    imp.plot(kind="barh", color="#7b2d43")
    plt.title(f"Feature Importance ({best_name})")
    plt.xlabel("Importance")
    plt.tight_layout()
    plt.savefig(f"{OUT}/06_feature_importance.png", dpi=150)
    plt.close()

    # Plot 7: model comparison
    fig, axes = plt.subplots(1, 2, figsize=(9, 4))
    names = list(results)
    axes[0].bar(names, [results[n]["R2"] for n in names], color=["#c98b9b", "#7b2d43"])
    axes[0].set_title("R2 Score (higher is better)")
    axes[1].bar(names, [results[n]["MAE"] for n in names], color=["#c98b9b", "#7b2d43"])
    axes[1].set_title("MAE (lower is better)")
    for ax, key in zip(axes, ["R2", "MAE"]):
        for i, n in enumerate(names):
            ax.text(i, results[n][key], f"{results[n][key]:.3f}", ha="center", va="bottom")
    plt.tight_layout()
    plt.savefig(f"{OUT}/07_model_comparison.png", dpi=150)
    plt.close()

    # --------------------------------------------------------------
    # SAVE MODEL + INFO FOR THE STREAMLIT APP
    # --------------------------------------------------------------
    joblib.dump(best_pipe, f"{OUT}/wine_model.pkl")
    feature_info = {
        "features": feature_names,
        "min": X.min().to_dict(),
        "max": X.max().to_dict(),
        "median": X.median().to_dict(),
        "importance": dict(zip(feature_names, map(float, importance))),
        "best_model": best_name,
    }
    with open(f"{OUT}/feature_info.json", "w") as f:
        json.dump(feature_info, f, indent=2)
    with open(f"{OUT}/metrics.json", "w") as f:
        json.dump(results, f, indent=2)

    # --------------------------------------------------------------
    # NUMBERS USED IN THE WRITTEN ANALYSIS
    # --------------------------------------------------------------
    corr_q = df.corr()["quality"].drop("quality").sort_values()
    top_pos, top_neg = corr_q.index[-1], corr_q.index[0]
    top3 = list(imp.sort_values(ascending=False).index[:3])
    share_56 = float(df["quality"].isin([5, 6]).mean())

    test_df = pd.DataFrame({"actual": np.asarray(y_test), "pred": best_pred})
    by_q = test_df.groupby("actual")["pred"].agg(["count", "mean"]).reset_index()
    high_mean = test_df.loc[test_df.actual >= 7, "pred"].mean()
    low_mean = test_df.loc[test_df.actual <= 4, "pred"].mean()

    other = [n for n in results if n != best_name][0]
    r_other = results[other]

    # --------------------------------------------------------------
    # 7. BUILD THE WORD REPORT
    # --------------------------------------------------------------
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Wine Quality Prediction using Machine Learning", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_par(doc, "Assignment 2 - Machine Learning Application", bold=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in [f"Name: {STUDENT_NAME}", f"Roll Number: {ROLL_NUMBER}", COLLEGE,
                 f"Course: {COURSE}", f"Date: {datetime.date.today():%d %B %Y}"]:
        add_par(doc, line).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Problem statement
    doc.add_heading("1. Problem Statement", level=1)
    add_par(doc,
        "Wine quality is normally judged by human experts, which is slow, subjective and "
        "expensive. The goal of this project is to build a machine learning model that "
        "predicts the quality score of a red wine using only measurable physicochemical "
        "properties such as acidity, sugar, sulphates and alcohol. Because the quality "
        "score is a number, this is a supervised regression problem.")
    add_par(doc, "Such a model can help producers screen batches quickly and understand which "
                 "chemical properties influence quality the most.")

    # 2. Dataset
    doc.add_heading("2. Dataset", level=1)
    add_par(doc,
        f"The Wine Quality dataset (red wine, Vinho Verde from Portugal) from the UCI Machine "
        f"Learning Repository was used (Cortez et al., 2009). It contains {n_rows_raw} samples and "
        f"{n_cols - 1} input features. The target is 'quality', a score given by wine experts.")
    add_par(doc, "Source: https://archive.ics.uci.edu/dataset/186/wine+quality", italic=True)
    add_table(doc, ["Feature", "Description"],
              [(c, FEATURE_DESCRIPTIONS.get(c, "")) for c in df.columns])
    add_fig(doc, f"{OUT}/01_quality_distribution.png",
            "Figure 1: Distribution of quality scores in the dataset")
    add_par(doc, f"Most wines are rated 5 or 6 ({share_56:.0%} of the data), so the data is "
                 f"imbalanced: very poor and excellent wines are rare.")

    # 3. Preprocessing
    doc.add_heading("3. Data Preprocessing", level=1)
    add_bullets(doc, [
        f"Missing values: {n_missing} found, so no imputation was needed."
        if n_missing == 0 else f"Missing values: {n_missing} found and filled with the column median.",
        f"Duplicate rows: {n_dupes} duplicates were removed, leaving {n_rows} unique samples. "
        f"This prevents identical wines appearing in both the training and test sets.",
        "Feature scaling: all features were standardised with StandardScaler (mean 0, standard "
        "deviation 1) because the features are on very different scales (e.g. total sulfur "
        "dioxide up to ~290 vs. chlorides around 0.08).",
        "Scaling was placed inside a scikit-learn Pipeline so that it is fitted only on the "
        "training data, which avoids data leakage.",
        f"Train/test split: 80% training ({len(X_train)} samples) and 20% testing ({len(X_test)} "
        f"samples), stratified by quality so both sets have similar class proportions.",
    ])
    add_fig(doc, f"{OUT}/02_correlation_heatmap.png", "Figure 2: Correlation between features and quality")
    add_fig(doc, f"{OUT}/03_alcohol_vs_quality.png", "Figure 3: Alcohol content increases with wine quality", width=5.0)

    # 4. Algorithm
    doc.add_heading("4. Algorithm Selection", level=1)
    add_par(doc, "Two regression algorithms were compared:")
    add_bullets(doc, [
        "Linear Regression - a simple baseline that assumes a straight-line relationship "
        "between the features and quality. It is fast and easy to interpret.",
        "Random Forest Regressor - an ensemble of 300 decision trees whose predictions are "
        "averaged. It captures non-linear relationships and interactions between features, "
        "handles outliers well, and gives feature importance scores.",
    ])
    add_par(doc, f"The final model was selected using 5-fold cross-validation R2 on the training "
                 f"data only. The selected model is: {best_name}.")

    # 5. Training and testing
    doc.add_heading("5. Model Training and Testing", level=1)
    add_par(doc,
        f"Both models were trained on {len(X_train)} samples and tested on the {len(X_test)} "
        f"unseen samples. Random state 42 was used so the results are reproducible. "
        f"5-fold cross-validation on the training set was also used to check that the results "
        f"are stable.")

    # 6. Metrics
    doc.add_heading("6. Evaluation Metrics", level=1)
    add_bullets(doc, [
        "R2 score - fraction of the variation in quality explained by the model (1.0 is perfect).",
        "MAE (Mean Absolute Error) - average size of the error in quality points.",
        "RMSE (Root Mean Squared Error) - like MAE but penalises large errors more.",
        "Exact accuracy - percentage of wines whose rounded prediction equals the true score.",
        "Within +/-1 accuracy - percentage of rounded predictions that are at most 1 point off.",
    ])
    add_table(doc,
        ["Model", "R2", "MAE", "RMSE", "Exact acc.", "Within 1 pt", "CV R2 (mean)"],
        [(n, f"{r['R2']:.3f}", f"{r['MAE']:.3f}", f"{r['RMSE']:.3f}",
          f"{r['Exact_Accuracy']:.1%}", f"{r['Within_1_Point']:.1%}", f"{r['CV_R2']:.3f}")
         for n, r in results.items()])

    # 7. Results
    doc.add_heading("7. Results and Graphs", level=1)
    add_fig(doc, f"{OUT}/07_model_comparison.png", "Figure 4: Model comparison", width=5.8)
    add_fig(doc, f"{OUT}/04_actual_vs_predicted.png", f"Figure 5: Actual vs predicted quality ({best_name})", width=4.8)
    add_fig(doc, f"{OUT}/05_residuals.png", "Figure 6: Distribution of prediction errors", width=5.0)
    add_fig(doc, f"{OUT}/06_feature_importance.png", f"Figure 7: Feature importance ({best_name})", width=5.2)
    add_par(doc, "Average prediction by true quality (test set):", bold=True)
    add_table(doc, ["True quality", "Samples", "Mean predicted quality"],
              [(int(r.actual), int(r["count"]), f"{r['mean']:.2f}") for _, r in by_q.iterrows()])

    # 8. Analysis
    doc.add_heading("8. Analysis of Results", level=1)
    better = best["R2"] >= r_other["R2"]
    add_par(doc,
        f"The {best_name} model achieved an R2 of {best['R2']:.3f} on the unseen test data, so it "
        f"explains about {best['R2']:.0%} of the variation in wine quality. Its MAE of "
        f"{best['MAE']:.2f} means a prediction is on average only {best['MAE']:.2f} points away "
        f"from the expert score, and {best['Within_1_Point']:.0%} of the rounded predictions are "
        f"within one point of the true score. The exact-match accuracy is "
        f"{best['Exact_Accuracy']:.0%}, which is lower because quality is scored in whole "
        f"numbers and neighbouring scores are hard to separate from chemistry alone.")
    add_par(doc,
        f"Compared with {other} (R2 = {r_other['R2']:.3f}, MAE = {r_other['MAE']:.2f}), "
        f"{best_name} {'performed better' if better else 'performed similarly'}. "
        + ("This suggests the relationship between chemical properties and quality is not purely "
           "linear, and the tree-based model captures interactions that a straight line misses."
           if best_name == "Random Forest" else
           "This suggests a mostly linear relationship between the chemical properties and quality."))
    add_par(doc,
        f"The most important features were {top3[0]}, {top3[1]} and {top3[2]}. In the correlation "
        f"analysis, {top_pos} had the strongest positive correlation with quality "
        f"({corr_q[top_pos]:+.2f}) and {top_neg} the strongest negative correlation "
        f"({corr_q[top_neg]:+.2f}). Figure 3 also shows that better wines tend to have higher alcohol.")
    lim = ("Limitations: the data is imbalanced, since most wines score 5 or 6. "
           "As a result the model tends to predict values near the middle. ")
    if not (np.isnan(high_mean) or np.isnan(low_mean)):
        lim += (f"On the test set, wines with true quality 7 or above were on average predicted as "
                f"{high_mean:.2f}, while wines with true quality 4 or below were predicted as "
                f"{low_mean:.2f}, so the extremes are pulled towards the average. ")
    lim += ("The dataset also covers only red Vinho Verde wine, and quality scores are subjective, "
            "so the model may not generalise to other wine types.")
    add_par(doc, lim)
    add_par(doc, "Possible improvements: hyperparameter tuning (GridSearchCV), trying Gradient "
                 "Boosting or XGBoost, treating the problem as classification (poor / average / good), "
                 "using oversampling techniques such as SMOTE for rare classes, and adding the white "
                 "wine data.")

    # 9. Streamlit
    doc.add_heading("9. Streamlit Interface (Optional Extension)", level=1)
    add_par(doc,
        "A simple web interface was built with Streamlit (app.py). The user adjusts sliders for "
        "the 11 chemical properties and clicks 'Predict Quality'. The saved model is loaded and "
        "returns the predicted quality score with a short interpretation.")

    # 10. Conclusion
    doc.add_heading("10. Conclusion", level=1)
    add_par(doc,
        f"A machine learning model was built to predict red wine quality from chemical "
        f"properties. After cleaning and scaling the data, {best_name} gave the best results "
        f"(R2 = {best['R2']:.3f}, MAE = {best['MAE']:.2f}). {top3[0].capitalize()}, {top3[1]} "
        f"and {top3[2]} are the main drivers of quality. Although the model cannot perfectly "
        f"reproduce expert scores, it provides a fast, cheap first estimate of wine quality.")

    doc.add_heading("References", level=1)
    add_bullets(doc, [
        "P. Cortez, A. Cerdeira, F. Almeida, T. Matos and J. Reis. Modeling wine preferences by "
        "data mining from physicochemical properties. Decision Support Systems, 47(4):547-553, 2009.",
        "UCI Machine Learning Repository - Wine Quality dataset: "
        "https://archive.ics.uci.edu/dataset/186/wine+quality",
        "Pedregosa et al. Scikit-learn: Machine Learning in Python. JMLR 12, 2011.",
    ])

    doc.save(f"{OUT}/Wine_Quality_Report.docx")
    print(f"\nDone. Check the '{OUT}' folder for plots, the model and Wine_Quality_Report.docx")


if __name__ == "__main__":
    main()
